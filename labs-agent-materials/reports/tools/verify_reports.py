# -*- coding: utf-8 -*-
"""Проверка отчётов: число страниц (Word), PDF, запрещённые слова, существование путей.

Запуск из корня репозитория:
    py labs-agent-materials\\reports\\tools\\verify_reports.py [каталог_для_pdf]
"""
import glob
import json
import os
import re
import sys
import tempfile

from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
REPORTS = os.path.join(ROOT, "labs-agent-materials", "reports")
FORBIDDEN = re.compile(
    r"llm|нейро|gpt|openai|openrouter|PROVIDED|BranchValueProvider|AgentAssembly|Suggester|языков[а-я]+ модел",
    re.IGNORECASE,
)
PATH_RE = re.compile(r"`([A-Za-z0-9_\-./]+\.(?:kt|py|json|sql|md|ps1|db|jar|puml|png))`")


def document_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in doc.sections:
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def mentioned_paths(text):
    found = set()
    for value in PATH_RE.findall(text):
        found.add(value)
    for value in re.findall(r"\(((?:src|labs-agent-materials)/[^,()]+?), строки", text):
        found.add(value)
    for value in re.findall(r"\((src/[^()]+?\.sql)\)", text):
        found.add(value)
    return sorted(found)


def resolve(value):
    candidates = [
        os.path.join(ROOT, value),
        os.path.join(ROOT, "src", "main", "kotlin", "com", "example", "agent", value),
        os.path.join(ROOT, "labs-agent-materials", "lab10-uml", value),
        os.path.join(ROOT, "labs-agent-materials", "lab11-database", value),
        os.path.join(ROOT, "labs-agent-materials", "lab12-app", value),
    ]
    if any(os.path.exists(c) for c in candidates):
        return True
    name = os.path.basename(value)
    return bool(glob.glob(os.path.join(ROOT, "src", "**", name), recursive=True)
                or glob.glob(os.path.join(ROOT, "labs-agent-materials", "**", name), recursive=True))


def word_pages(paths, pdf_dir):
    import win32com.client
    result = {}
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        for path in paths:
            doc = word.Documents.Open(os.path.abspath(path), False, True)
            try:
                doc.Repaginate()
                pages = doc.ComputeStatistics(2)
                pdf = os.path.join(pdf_dir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
                doc.SaveAs2(pdf, FileFormat=17)
                result[os.path.basename(path)] = {"pages": pages, "pdf": pdf}
            finally:
                doc.Close(False)
    finally:
        word.Quit()
    return result


def blank_pages(pdf):
    import fitz
    blanks = []
    with fitz.open(pdf) as doc:
        for index, page in enumerate(doc, 1):
            text = page.get_text().strip()
            images = page.get_images()
            if len(text) < 15 and not images:
                blanks.append(index)
    return blanks


def main():
    pdf_dir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(tempfile.gettempdir(), "lab-reports-check")
    os.makedirs(pdf_dir, exist_ok=True)
    reports = sorted(glob.glob(os.path.join(REPORTS, "*.docx")))
    summary = {}
    for path in reports:
        text = document_text(path)
        hits = sorted({m.group(0) for m in FORBIDDEN.finditer(text)})
        paths = mentioned_paths(text)
        missing = [p for p in paths if not resolve(p)]
        summary[os.path.basename(path)] = {"forbidden": hits, "paths": len(paths), "missing_paths": missing}
    for name, info in word_pages(reports, pdf_dir).items():
        summary[name].update(info)
        summary[name]["blank_pages"] = blank_pages(info["pdf"])
    out = os.path.join(pdf_dir, "summary.json")
    with open(out, "w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=1)
    print(out)


if __name__ == "__main__":
    main()
