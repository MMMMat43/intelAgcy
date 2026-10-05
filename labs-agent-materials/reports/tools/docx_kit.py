# -*- coding: utf-8 -*-
"""Общие функции оформления отчётов по требованиям ВятГУ (python-docx)."""
import os
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
CODE_FONT = "Consolas"
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))


def path(*parts):
    return os.path.join(ROOT, *parts)


def _font(run, size=14, bold=False, italic=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), name)


def _margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)


def _page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run()
    _font(run, 12)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._element.append(element)


class Report:
    def __init__(self):
        self.doc = Document()
        style = self.doc.styles["Normal"]
        style.font.name = FONT
        style.font.size = Pt(14)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        _margins(self.doc.sections[0])
        self.figure = 0
        self.table_no = 0
        self.listing = 0

    def _paragraph(self, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, before=0, after=0, spacing=1.5, keep=False):
        p = self.doc.add_paragraph()
        p.alignment = align
        f = p.paragraph_format
        f.first_line_indent = Cm(1.25) if indent else Cm(0)
        f.space_before = Pt(before)
        f.space_after = Pt(after)
        f.line_spacing = spacing
        f.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        if keep:
            f.keep_with_next = True
        return p

    def title_page(self, number, title):
        lines = [
            ("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ", False),
            ("Федеральное государственное бюджетное образовательное учреждение высшего образования", False),
            ("«ВЯТСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ»", True),
            ("Институт математики и информационных систем", False),
            ("Факультет автоматики и вычислительной техники", False),
            ("Кафедра электронных вычислительных машин", False),
        ]
        for text, bold in lines:
            p = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.15)
            _font(p.add_run(text), 14, bold=bold)
        for _ in range(5):
            self._paragraph(indent=False, spacing=1.15)
        p = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.15)
        _font(p.add_run(f"ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ №{number}"), 16, bold=True)
        p = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=6, spacing=1.15)
        _font(p.add_run(f"«{title}»"), 16, bold=True)
        for _ in range(6):
            self._paragraph(indent=False, spacing=1.15)
        for text in (
            "Выполнил: студент группы _________",
            "Кожин М. С.  _____________",
            "",
            "Проверил: _________________",
            "",
            "Оценка: ________  Дата: ________",
        ):
            p = self._paragraph(WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=1.15)
            p.paragraph_format.left_indent = Cm(8.0)
            _font(p.add_run(text), 14)
        for _ in range(5):
            self._paragraph(indent=False, spacing=1.15)
        p = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, indent=False, spacing=1.15)
        _font(p.add_run("Киров 2026"), 14)
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        _margins(section)
        _page_number_footer(section)
        title_footer = self.doc.sections[0].footer
        title_footer.is_linked_to_previous = False

    def heading(self, text, level=1):
        p = self._paragraph(WD_ALIGN_PARAGRAPH.LEFT if level > 1 else WD_ALIGN_PARAGRAPH.LEFT,
                            indent=True, before=12 if level == 1 else 8, after=6, keep=True)
        _font(p.add_run(text), 14, bold=True)

    def text(self, value):
        for chunk in [part.strip() for part in value.strip().split("\n\n") if part.strip()]:
            p = self._paragraph()
            self._inline(p, " ".join(line.strip() for line in chunk.splitlines()))

    def _inline(self, paragraph, value):
        parts = value.split("`")
        for index, part in enumerate(parts):
            if not part:
                continue
            run = paragraph.add_run(part)
            if index % 2 == 1:
                _font(run, 12, name=CODE_FONT)
            else:
                _font(run, 14)

    def bullets(self, items, numbered=False):
        for index, item in enumerate(items, 1):
            p = self._paragraph()
            marker = f"{index}) " if numbered else "– "
            self._inline(p, marker + item)

    def figure_image(self, image, caption, width_cm=16.0):
        p = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=6, spacing=1.0, keep=True)
        p.add_run().add_picture(image, width=Cm(width_cm))
        self.figure += 1
        c = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6, spacing=1.0)
        _font(c.add_run(f"Рисунок {self.figure} – {caption}"), 12)
        return self.figure

    def next_figure(self):
        return self.figure + 1

    def next_table(self):
        return self.table_no + 1

    def table(self, caption, header, rows, widths=None, size=11):
        self.table_no += 1
        c = self._paragraph(WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=6, after=3, spacing=1.0, keep=True)
        _font(c.add_run(f"Таблица {self.table_no} – {caption}"), 12)
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, value in zip(table.rows[0].cells, header):
            cell.text = ""
            _font(cell.paragraphs[0].add_run(value), size, bold=True)
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.line_spacing = 1.0
                text = str(value)
                parts = text.split("`")
                for index, part in enumerate(parts):
                    if part:
                        _font(paragraph.add_run(part), size - 1 if index % 2 else size,
                              name=CODE_FONT if index % 2 else FONT)
        if widths:
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        self._paragraph(indent=False, spacing=1.0)
        return self.table_no

    def code(self, caption, source_lines, size=9):
        self.listing += 1
        c = self._paragraph(WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=6, after=3, spacing=1.0, keep=True)
        _font(c.add_run(f"Листинг {self.listing} – {caption}"), 12)
        for index, line in enumerate(source_lines):
            p = self._paragraph(WD_ALIGN_PARAGRAPH.LEFT, indent=False, spacing=1.0)
            p.paragraph_format.left_indent = Cm(0.3)
            if index < len(source_lines) - 1:
                p.paragraph_format.keep_with_next = False
            _font(p.add_run(line.rstrip("\r") if line.strip() else " "), size, name=CODE_FONT)
        self._paragraph(indent=False, spacing=1.0)
        return self.listing

    def landscape(self):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(1.5)
        _page_number_footer(section)

    def portrait(self):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.PORTRAIT
        _margins(section)
        _page_number_footer(section)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def save(self, target):
        os.makedirs(os.path.dirname(target), exist_ok=True)
        self.doc.save(target)


def read_lines(relative, start=None, end=None, skip=None):
    with open(path(*relative.split("/")), encoding="utf-8") as handle:
        lines = handle.read().splitlines()
    if start is not None:
        lines = lines[start - 1:end]
    if skip:
        lines = [line for line in lines if not any(token in line for token in skip)]
    return lines
